from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = r'D:\GameDev\-blender-3d'
TEMPLATE = r'F:\_แบบฝึกหัดที่ 6  ออกแบบตัวละคร.docx'
OUTPUT = ROOT + r'\แบบฝึกหัดที่ 6 ออกแบบตัวละคร ฉบับกรอกข้อมูล.docx'

doc = Document(TEMPLATE)
body = doc._element.body
for element in list(body):
    if element.tag != qn('w:sectPr'):
        body.remove(element)

section = doc.sections[0]
section.top_margin = Inches(.68)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.73)
section.right_margin = Inches(.73)

normal = doc.styles['Normal']
normal.font.name = 'TH Sarabun New'
normal.font.size = Pt(14)
normal.font.color.rgb = RGBColor(30, 30, 30)
normal.paragraph_format.space_after = Pt(7)
for name in ('Title', 'Heading 1'):
    s = doc.styles[name]
    s.font.name = 'TH Sarabun New'
    s.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)

def paragraph(text='', style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def picture(path, width, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(path, width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(11)
    for run in c.runs:
        run.italic = True
        run.font.size = Pt(11)

def link(label, url):
    p = doc.add_paragraph()
    p.add_run(label + '  ')
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1765A1')
    props.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    props.append(underline)
    r.append(props)
    t = OxmlElement('w:t')
    t.text = url
    r.append(t)
    h.append(r)
    p._p.append(h)

title = paragraph('แบบฝึกหัดที่ 6 ออกแบบตัวละคร 3D', 'Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph('ผู้จัดทำ: Natdanai137  |  Blender 4.2 และ Godot 4.7')
paragraph('สร้างตัวละคร Mira แบบ low poly ใน Blender พร้อมโครงกระดูก humanoid แล้วนำเข้าฉากสาธิตใน Godot เพื่อแสดงท่าทางจากคลัง Mixamo ที่เตรียมไว้ และส่งออกเป็น Web ให้ทดลองเล่นได้')

paragraph('1. ภาพโมเดลจาก Blender3D', 'Heading 1')
paragraph('ตัวละคร Mira สร้างจากชิ้นส่วนทรงพื้นฐานใน Blender ใส่วัสดุสีและผูกกับ armature ชื่อกระดูกแบบ Mixamo บันทึกต้นฉบับเป็น source/Mira.blend และส่งออกเป็น assets/Mira.glb')
picture(ROOT + r'\screenshots\Blender-Mira.png', 3.55, 'ภาพที่ 1  โมเดล Mira ที่สร้างและเรนเดอร์จาก Blender')

doc.add_page_break()
paragraph('2. ท่าทางจาก Mixamo Animation Library', 'Heading 1')
paragraph('นำ Mixamo BoneMap.tres ที่เตรียมไว้มาใช้กับโครงกระดูกของ Mira ใน Godot แล้วโหลด MeleeLib.res และ ShooterLib.res เข้ากับ AnimationPlayer ตัวอย่างนี้แสดงท่า Slash1 จากคลัง Melee')
picture(ROOT + r'\screenshots\Godot-slash.png', 6.55, 'ภาพที่ 2  ตัวละคร Mira แสดงท่า Slash1 จากไลบรารีแอนิเมชัน')
paragraph('งานนี้ใช้ไฟล์แอนิเมชันที่จัดเตรียมไว้ในโปรเจกต์ จึงไม่มีภาพขั้นตอนอัปโหลดตัวละครผ่านเว็บไซต์ Mixamo โดยตรง')

doc.add_page_break()
paragraph('3. ฉาก Demo ใน Godot', 'Heading 1')
paragraph('ฉากเดโมมีพื้นหญ้า ต้นไม้ low poly และโมเดล Mako จาก Poly Pizza ประกอบฉาก สามารถกด Idle, Walk, Run, Slash, Jump, Aim และ Punch เพื่อเปลี่ยนท่าตัวละคร รวมทั้งลากเมาส์เพื่อหมุนกล้องและเลื่อนล้อเมาส์เพื่อซูม')
picture(ROOT + r'\screenshots\Godot-demo.png', 6.55, 'ภาพที่ 3  ฉากเดโมใน Godot Web export พร้อมตัวละคร Mira และ Mako')

paragraph('4. ลิงก์ส่งงาน', 'Heading 1')
link('a. ทดลองเล่นบนเว็บ', 'https://natdanai137.github.io/Blender3D/')
link('b. โครงการ GitHub', 'https://github.com/Natdanai137/Blender3D')
paragraph('ที่มาโมเดลประกอบ: Mako โดย Quaternius บน Poly Pizza (CC0) — https://poly.pizza/m/2urczqZ9Xf')

doc.save(OUTPUT)
print(OUTPUT)
